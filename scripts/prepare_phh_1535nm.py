"""Prepare the Figure 4(b) COMSOL sweep for PhotonAct without modifying the source workbook."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from collections.abc import Sequence
from math import isfinite
from pathlib import Path
from typing import Any


def _read_branch(sheet: Any, x_column: int, y_column: int) -> list[tuple[float, float]]:
    points: list[tuple[float, float]] = []
    for row in range(1, sheet.max_row + 1):
        x = sheet.cell(row, x_column).value
        y = sheet.cell(row, y_column).value
        if x is None and y is None:
            continue
        if isinstance(x, str) and y is None:
            continue
        if not isinstance(x, int | float) or not isinstance(y, int | float):
            raise ValueError(f"Expected numeric values at row {row}")
        point = (float(x), float(y))
        if not all(isfinite(value) for value in point):
            raise ValueError(f"Non-finite value at row {row}")
        points.append(point)
    points.sort(key=lambda point: point[0])
    if len(points) < 2 or any(
        a[0] >= b[0] for a, b in zip(points, points[1:], strict=False)
    ):
        raise ValueError("Each branch needs strictly increasing, unique input values")
    return points


def _transition_midpoint(points: list[tuple[float, float]]) -> tuple[float, tuple[float, float]]:
    left, right = max(
        zip(points, points[1:], strict=False),
        key=lambda pair: abs(pair[1][1] - pair[0][1]),
    )
    return (left[0] + right[0]) / 2.0, (left[0], right[0])


def prepare(source: Path, output_dir: Path, data_license: str) -> tuple[Path, Path]:
    """Extract fine up/down sweeps and write canonical CSV plus provenance metadata."""
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError("Install PhotonAct with the 'data' extra to read .xlsx files") from error

    workbook = load_workbook(source, data_only=True, read_only=True)
    if "Sheet1" not in workbook.sheetnames:
        raise ValueError("Expected a Sheet1 worksheet")
    sheet = workbook["Sheet1"]
    up = _read_branch(sheet, 8, 9)  # H:I, low-to-high scan
    down = _read_branch(sheet, 11, 12)  # K:L, high-to-low scan; sorted here
    if [point[0] for point in up] != [point[0] for point in down]:
        raise ValueError("Up/down sweeps must use the same input grid")

    upper_threshold, upper_interval = _transition_midpoint(up)
    lower_threshold, lower_interval = _transition_midpoint(down)
    if lower_threshold >= upper_threshold:
        raise ValueError("Detected thresholds do not form a hysteresis interval")

    output_dir.mkdir(parents=True, exist_ok=True)
    csv_path = output_dir / "phh_1535nm.csv"
    metadata_path = output_dir / "phh_1535nm.meta.json"
    with csv_path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(
            stream,
            fieldnames=["input_power", "transmittance", "output_power", "branch"],
        )
        writer.writeheader()
        for branch, points in (("up", up), ("down", down)):
            for input_power, transmittance in points:
                writer.writerow(
                    {
                        "input_power": format(input_power, ".12g"),
                        "transmittance": format(transmittance, ".12g"),
                        "output_power": format(input_power * transmittance, ".12g"),
                        "branch": branch,
                    }
                )

    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    metadata = {
        "name": "phh_1535nm",
        "input_unit": "W/m",
        "output_unit": "W/m",
        "source": "COMSOL simulation data underlying Figure 4(b), Sheet1 H1:I1000 and K1:L1000",
        "valid_min": up[0][0],
        "valid_max": up[-1][0],
        "normalized": False,
        "license": data_license,
        "citation": "https://doi.org/10.3390/nano16090561",
        "notes": (
            "Output power is computed as input power multiplied by power transmittance. "
            f"Source SHA-256: {digest}. Observed transition intervals: "
            f"down {lower_interval[0]:g}-{lower_interval[1]:g} W/m; "
            f"up {upper_interval[0]:g}-{upper_interval[1]:g} W/m. "
            "Thresholds are the interval midpoints."
        ),
        "data_kind": "simulated",
        "response_quantity": "output_power_from_power_transmittance",
        "wavelength_nm": 1535.0,
        "polarization": "TM",
        "lower_threshold": lower_threshold,
        "upper_threshold": upper_threshold,
    }
    metadata_path.write_text(json.dumps(metadata, indent=2) + "\n", encoding="utf-8")
    return csv_path, metadata_path


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="Path to shuangwentiai.xlsx")
    parser.add_argument(
        "--output-dir", type=Path, default=Path("local_data/phh_1535nm")
    )
    parser.add_argument(
        "--data-license",
        default="All rights reserved; publication approval pending",
        help="License recorded in the generated metadata",
    )
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    """Run the deterministic workbook-to-curve conversion."""
    arguments = _parser().parse_args(argv)
    csv_path, metadata_path = prepare(
        arguments.source, arguments.output_dir, arguments.data_license
    )
    print(csv_path)
    print(metadata_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
