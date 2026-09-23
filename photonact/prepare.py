"""Prepare external curve tables without changing the source file."""

from __future__ import annotations

import csv
import hashlib
import json
import re
import tempfile
from dataclasses import asdict, dataclass
from math import isfinite
from pathlib import Path
from typing import Any, Literal

from photonact.curves import load_curve

Transform = Literal["direct", "input-times-transmittance"]


@dataclass(frozen=True)
class PrepareOptions:
    """Explicit column mapping and provenance for a local curve conversion."""

    name: str
    input_unit: str
    output_unit: str
    source_description: str
    data_kind: str
    license: str
    input_column: str | None = None
    output_column: str | None = None
    branch_column: str | None = None
    up_input_column: str | None = None
    up_output_column: str | None = None
    down_input_column: str | None = None
    down_output_column: str | None = None
    sheet: str | None = None
    start_row: int = 1
    end_row: int | None = None
    sort: bool = False
    transform: Transform = "direct"
    citation: str = ""
    wavelength_nm: float | None = None
    polarization: str = ""
    lower_threshold: float | None = None
    upper_threshold: float | None = None

    def __post_init__(self) -> None:
        """Reject ambiguous mappings before reading a source file."""
        if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]*", self.name):
            raise ValueError("Name must use letters, digits, dots, underscores or hyphens")
        long_mapping = self.input_column is not None or self.output_column is not None
        wide = (
            self.up_input_column,
            self.up_output_column,
            self.down_input_column,
            self.down_output_column,
        )
        wide_mapping = any(value is not None for value in wide)
        if long_mapping == wide_mapping:
            raise ValueError("Choose either long or up/down column mapping")
        if long_mapping and (self.input_column is None or self.output_column is None):
            raise ValueError("Long mapping requires input and output columns")
        if wide_mapping and (not all(wide) or self.branch_column is not None):
            raise ValueError("Wide mapping requires four up/down columns and no branch column")
        if self.start_row < 1 or self.end_row is not None and self.end_row < self.start_row:
            raise ValueError("Invalid row range")
        if self.transform not in {"direct", "input-times-transmittance"}:
            raise ValueError("Unsupported output transformation")
        if self.data_kind not in {"measured", "simulated", "digitized", "synthetic"}:
            raise ValueError("Data kind must be measured, simulated, digitized or synthetic")
        if not self.input_unit.strip() or not self.output_unit.strip():
            raise ValueError("Input and output units must be stated")
        if not self.source_description.strip() or not self.license.strip():
            raise ValueError("Source and license must be stated")
        if (self.lower_threshold is None) != (self.upper_threshold is None):
            raise ValueError("Provide both hysteresis thresholds or neither")
        for value in (self.wavelength_nm, self.lower_threshold, self.upper_threshold):
            if value is not None and not isfinite(value):
                raise ValueError("Metadata values must be finite")
        if self.wavelength_nm is not None and self.wavelength_nm <= 0:
            raise ValueError("Wavelength must be positive")


def _column_index(column: str) -> int:
    """Convert a spreadsheet column letter or 1-based number to a zero-based index."""
    if column.isdecimal():
        index = int(column) - 1
    elif re.fullmatch(r"[A-Za-z]+", column):
        index = 0
        for letter in column.upper():
            index = index * 26 + ord(letter) - ord("A") + 1
        index -= 1
    else:
        raise ValueError(f"Invalid spreadsheet column: {column!r}")
    if index < 0:
        raise ValueError(f"Invalid spreadsheet column: {column!r}")
    return index


def _read_rows(source: Path, options: PrepareOptions) -> list[tuple[int, Any]]:
    """Read source rows with their original row numbers."""
    if source.suffix.lower() == ".csv":
        if options.sheet is not None:
            raise ValueError("--sheet is only valid for XLSX")
        with source.open(newline="", encoding="utf-8-sig") as stream:
            reader = csv.DictReader(stream)
            if not reader.fieldnames:
                raise ValueError("CSV needs a header row")
            return [
                (row_number, row)
                for row_number, row in enumerate(reader, 2)
                if row_number >= options.start_row
                and (options.end_row is None or row_number <= options.end_row)
            ]
    if source.suffix.lower() == ".xlsx":
        if options.sheet is None:
            raise ValueError("XLSX requires an explicit --sheet")
        try:
            from openpyxl import load_workbook  # type: ignore
        except ImportError as error:
            raise RuntimeError('Install the XLSX extra with: pip install -e ".[data]"') from error
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            if options.sheet not in workbook.sheetnames:
                raise ValueError(f"Worksheet not found: {options.sheet}")
            sheet = workbook[options.sheet]
            return [
                (row_number, tuple(row))
                for row_number, row in enumerate(
                    sheet.iter_rows(
                        min_row=options.start_row,
                        max_row=options.end_row,
                        values_only=True,
                    ),
                    options.start_row,
                )
            ]
        finally:
            workbook.close()
    raise ValueError("Source must be .csv or .xlsx")


def _cell(row: Any, column: str, row_number: int) -> Any:
    if isinstance(row, dict):
        if column not in row:
            raise ValueError(f"CSV column not found: {column!r}")
        return row[column]
    index = _column_index(column)
    if index >= len(row):
        raise ValueError(f"Column {column!r} is missing at row {row_number}")
    return row[index]


def _numeric(value: Any, row_number: int, column: str) -> float:
    try:
        result = float(value)
    except (TypeError, ValueError) as error:
        raise ValueError(f"Expected a number at row {row_number}, column {column}") from error
    if not isfinite(result):
        raise ValueError(f"Non-finite value at row {row_number}, column {column}")
    return result


def _collect_points(
    rows: list[tuple[int, Any]], options: PrepareOptions
) -> list[dict[str, float | str]]:
    branches: dict[str, list[tuple[int, float, float]]] = {}
    for row_number, row in rows:
        if options.input_column is not None and options.output_column is not None:
            mappings = [(
                options.branch_column,
                options.input_column,
                options.output_column,
            )]
        else:
            assert options.up_input_column and options.up_output_column
            assert options.down_input_column and options.down_output_column
            mappings = [
                ("up", options.up_input_column, options.up_output_column),
                ("down", options.down_input_column, options.down_output_column),
            ]
        for branch_source, x_column, y_column in mappings:
            x_value = _cell(row, x_column, row_number)
            y_value = _cell(row, y_column, row_number)
            if x_value in (None, "") and y_value in (None, ""):
                continue
            if options.branch_column is not None:
                branch_value = _cell(row, options.branch_column, row_number)
                if branch_value in (None, ""):
                    raise ValueError(f"Missing branch at row {row_number}")
                branch = str(branch_value).strip().lower()
            else:
                branch = branch_source or "single"
            if not branch:
                raise ValueError(f"Empty branch at row {row_number}")
            x = _numeric(x_value, row_number, x_column)
            y = _numeric(y_value, row_number, y_column)
            branches.setdefault(branch, []).append((row_number, x, y))
    if not branches:
        raise ValueError("No numeric curve rows found")
    if options.lower_threshold is not None and set(branches) != {"up", "down"}:
        raise ValueError("Hysteresis thresholds require up and down branches")
    points: list[dict[str, float | str]] = []
    for branch, values in branches.items():
        if options.sort:
            values.sort(key=lambda point: point[1])
        if len(values) < 2:
            raise ValueError(f"Branch {branch!r} needs at least two points")
        for previous, current in zip(values, values[1:], strict=False):
            if previous[1] >= current[1]:
                raise ValueError(
                    f"Input in branch {branch!r} must increase at row {current[0]} "
                    "(use --sort only if source order is not meaningful)"
                )
        for _, x, y in values:
            output = x * y if options.transform == "input-times-transmittance" else y
            if not isfinite(output):
                raise ValueError("Transformed output must be finite")
            point: dict[str, float | str] = {
                "input_power": x,
                "output_power": output,
                "branch": branch,
            }
            if options.transform == "input-times-transmittance":
                point["transmittance"] = y
            points.append(point)
    if options.lower_threshold is not None:
        up = branches["up"]
        down = branches["down"]
        if (up[0][1], up[-1][1]) != (down[0][1], down[-1][1]):
            raise ValueError("Hysteresis branches must have the same input range")
    return points


def prepare_curve(
    source: str | Path, output_dir: str | Path, options: PrepareOptions
) -> tuple[Path, Path, Path]:
    """Convert a local source table to canonical curve, metadata, and audit report."""
    source_path = Path(source)
    if not source_path.is_file():
        raise FileNotFoundError(source_path)
    output_path = Path(output_dir)
    csv_path = output_path / f"{options.name}.csv"
    metadata_path = output_path / f"{options.name}.meta.json"
    report_path = output_path / f"{options.name}.preparation.json"
    for path in (csv_path, metadata_path, report_path):
        if path.exists():
            raise FileExistsError(f"Output already exists: {path}")
    points = _collect_points(_read_rows(source_path, options), options)
    x_values = [float(point["input_power"]) for point in points]
    metadata = {
        "name": options.name,
        "input_unit": options.input_unit,
        "output_unit": options.output_unit,
        "source": options.source_description,
        "valid_min": min(x_values),
        "valid_max": max(x_values),
        "license": options.license,
        "citation": options.citation,
        "data_kind": options.data_kind,
        "response_quantity": (
            "output_power_from_power_transmittance"
            if options.transform == "input-times-transmittance"
            else "output_power"
        ),
        "wavelength_nm": options.wavelength_nm,
        "polarization": options.polarization,
        "lower_threshold": options.lower_threshold,
        "upper_threshold": options.upper_threshold,
    }
    output_path.mkdir(parents=True, exist_ok=True)
    fieldnames = ["input_power", "output_power", "branch"]
    if options.transform == "input-times-transmittance":
        fieldnames.insert(1, "transmittance")
    digest = hashlib.sha256(source_path.read_bytes()).hexdigest()
    report = {
        "source_file": str(source_path.resolve()),
        "source_sha256": digest,
        "options": asdict(options),
        "point_count": len(points),
        "branches": sorted({str(point["branch"]) for point in points}),
    }
    with tempfile.TemporaryDirectory(prefix=".photonact-", dir=output_path) as stage_dir:
        stage = Path(stage_dir)
        staged_csv = stage / csv_path.name
        staged_metadata = stage / metadata_path.name
        staged_report = stage / report_path.name
        with staged_csv.open("w", newline="", encoding="utf-8") as stream:
            writer = csv.DictWriter(stream, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(points)
        staged_metadata.write_text(json.dumps(metadata, indent=2) + "\n", encoding="utf-8")
        load_curve(staged_csv)
        staged_report.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
        for staged, final in (
            (staged_csv, csv_path),
            (staged_metadata, metadata_path),
            (staged_report, report_path),
        ):
            if final.exists():
                raise FileExistsError(f"Output already exists: {final}")
            staged.rename(final)
    return csv_path, metadata_path, report_path
